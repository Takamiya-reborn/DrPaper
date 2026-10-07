"""docx 导出：把草稿节点渲染为规范排版的 Word 文档。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table as DocxTable

from drpaper.export.guides import guide_for_heading
from drpaper.export.math import latex_to_omml
from drpaper.export.style_profile import StyleProfile, load_profile
from drpaper.literature.base import Paper
from drpaper.paper.markdown_parser import (
    HeadingNode,
    MathNode,
    Node,
    ParagraphNode,
    TableNode,
    TitleNode,
    parse,
)

REFERENCES_HEADING = "参考文献"

ALIGN_JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
ALIGN_CENTER = WD_ALIGN_PARAGRAPH.CENTER


def export_docx(
    nodes: list[Node],
    references: list[Paper],
    path: str | Path,
    profile: StyleProfile | None = None,
) -> Path:
    """按样式档案渲染草稿与文献库，写出 docx 文件；未指定档案时用默认规范。"""
    profile = profile or load_profile()
    document = Document()
    _setup_page(document, profile)

    for node in nodes:
        if isinstance(node, TitleNode):
            _add_title(document, node.text, profile)
        elif isinstance(node, HeadingNode):
            _add_heading(document, node, profile)
        elif isinstance(node, TableNode):
            _add_table(document, node, profile)
        elif isinstance(node, MathNode):
            _add_display_math(document, node.text, profile)
        elif isinstance(node, ParagraphNode):
            _add_paragraph(document, node.runs, profile)

    if references:
        _add_references(document, references, profile)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    return out


def _setup_page(document: Document, profile: StyleProfile) -> None:
    margin = Cm(profile.margin_cm)
    for section in document.sections:
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin


def _set_fonts(run, profile: StyleProfile, cn_font: str, size_pt: float, bold: bool = False) -> None:
    """同时设置中英文字体（eastAsia 属性需手动写入 XML）。"""
    run.font.name = profile.en_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)


def _add_title(document: Document, text: str, profile: StyleProfile) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = ALIGN_CENTER
    paragraph.paragraph_format.line_spacing = profile.line_spacing
    _set_fonts(paragraph.add_run(text), profile, profile.cn_heading_font, profile.title_pt, bold=True)


def _add_heading(document: Document, node: HeadingNode, profile: StyleProfile) -> None:
    size_pt = profile.heading1_pt if node.level == 1 else profile.body_pt
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = profile.line_spacing
    paragraph.paragraph_format.space_before = Cm(0.3)
    paragraph.paragraph_format.space_after = Cm(0.2)
    _set_fonts(paragraph.add_run(node.text), profile, profile.cn_heading_font, size_pt, bold=True)

    guide = guide_for_heading(node.text)
    if guide:
        _attach_comment(document, paragraph, guide)


def _add_paragraph(document: Document, runs: list, profile: StyleProfile) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = ALIGN_JUSTIFY
    paragraph.paragraph_format.line_spacing = profile.line_spacing
    paragraph.paragraph_format.first_line_indent = Pt(
        profile.body_pt * profile.first_line_indent_chars
    )
    for run in runs:
        if run.math:
            _add_math_run(paragraph, run.text)
        else:
            _set_fonts(
                paragraph.add_run(run.text),
                profile,
                profile.cn_body_font,
                profile.body_pt,
                bold=run.bold,
            )


def _add_display_math(document: Document, latex: str, profile: StyleProfile) -> None:
    """渲染独立公式块：单独一段，居中。"""
    paragraph = document.add_paragraph()
    paragraph.alignment = ALIGN_CENTER
    paragraph.paragraph_format.line_spacing = profile.line_spacing
    _add_math_run(paragraph, latex)


def _add_math_run(paragraph, latex: str) -> None:
    """在段落中插入 Word 原生公式；转换失败时降级为斜体 LaTeX 原文。"""
    element = latex_to_omml(latex)
    if element is None:
        run = paragraph.add_run(latex)
        run.italic = True
        return
    paragraph._p.append(element)


def _add_table(document: Document, node: TableNode, profile: StyleProfile) -> None:
    """渲染学术三线表：表题居中加粗，顶线/底线 1.5pt，表头下线 0.75pt，无竖线。

    含†的预期值单元格加黄色高亮，并在表题上批注提醒学生替换为实测值。
    """
    caption_paragraph = None
    if node.caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = ALIGN_CENTER
        caption_paragraph.paragraph_format.line_spacing = profile.line_spacing
        _set_fonts(
            caption_paragraph.add_run(node.caption),
            profile,
            profile.cn_body_font,
            profile.table_pt,
            bold=True,
        )
    table = document.add_table(rows=len(node.rows) + 1, cols=len(node.header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    for col, text in enumerate(node.header):
        _set_cell(table.cell(0, col), text, profile, bold=True)
    for row, cells in enumerate(node.rows, start=1):
        for col, text in enumerate(cells):
            if col < len(node.header):
                _set_cell(table.cell(row, col), text, profile, highlight="†" in text)
    if caption_paragraph and any("†" in c for row in node.rows for c in row):
        _attach_comment(
            document,
            caption_paragraph,
            "标†的数值为基于文献库统计的预期参考区间，实验完成后请替换为实测值。",
        )


def _set_table_borders(table: DocxTable) -> None:
    """写 tblBorders XML：仅顶线与底线（三线表外框），竖线一律不加。"""
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = table._tbl.tblPr.makeelement(qn("w:tblBorders"), {})
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = borders.makeelement(qn(f"w:{edge}"), {})
            borders.append(element)
        visible = edge in ("top", "bottom")
        element.set(qn("w:val"), "single" if visible else "none")
        if visible:
            element.set(qn("w:sz"), "12")  # 1/8 pt 为单位，12 = 1.5pt
            element.set(qn("w:color"), "000000")
    # 表头行下线：逐单元格写 tcBorders
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.makeelement(qn("w:tcBorders"), {})
        bottom = tc_borders.makeelement(qn("w:bottom"), {})
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")  # 0.75pt
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)
        tc_pr.append(tc_borders)


def _set_cell(
    cell, text: str, profile: StyleProfile, bold: bool = False, highlight: bool = False
) -> None:
    """填充单元格：水平居中、表格字号、统一中英文字体；highlight 标记预期值。"""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = ALIGN_CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _set_fonts(run, profile, profile.cn_body_font, profile.table_pt, bold=bold)
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_references(document: Document, references: list[Paper], profile: StyleProfile) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.line_spacing = profile.line_spacing
    heading.paragraph_format.space_before = Cm(0.4)
    _set_fonts(heading.add_run(REFERENCES_HEADING), profile, profile.cn_heading_font, profile.heading1_pt, bold=True)

    for index, paper in enumerate(references, start=1):
        paragraph = document.add_paragraph()
        paragraph.alignment = ALIGN_JUSTIFY
        paragraph.paragraph_format.line_spacing = profile.line_spacing
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)  # 悬挂缩进
        _set_fonts(paragraph.add_run(paper.reference_line(index)), profile, profile.cn_body_font, profile.reference_pt)


def _attach_comment(document: Document, paragraph, text: str) -> None:
    """在节标题段落上附加一条写作指导批注（python-docx >= 1.2）。"""
    if paragraph.runs:
        document.add_comment(paragraph.runs[0], text=text, author="DrPaper", initials="DP")
